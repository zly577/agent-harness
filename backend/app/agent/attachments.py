"""Bounded attachment decoding and provider-independent message construction."""
import base64
import binascii
import io
import zipfile
from pathlib import Path

from fastapi import HTTPException

MAX_FILE_BYTES = 5 * 1024 * 1024
MAX_TOTAL_BYTES = 15 * 1024 * 1024
MAX_TEXT_CHARS = 100_000
REFUSAL = "我不是多模态模型，无法识别图片。请切换到支持图片识别的多模态模型后重试。"
TEXT_EXTENSIONS = {".txt", ".md", ".csv", ".tsv", ".json", ".xml", ".yaml", ".yml", ".log", ".py", ".js", ".ts", ".html", ".css"}
IMAGE_EXTENSIONS = {".png", ".jpg", ".jpeg", ".webp", ".gif"}


def supports_vision(cfg):
    explicit = cfg.parameters.get("supports_vision")
    if isinstance(explicit, bool):
        return explicit
    model = cfg.model.lower()
    # 未知/自部署模型需在设置中明确声明能力。
    return model.startswith(("gpt-4o", "gpt-4.1", "gpt-5", "claude-3", "claude-sonnet-4", "claude-opus-4", "claude-haiku-4"))


def prepare_content(req):
    if not req.attachments:
        return req.message, False
    blocks = [{"type": "text", "text": req.message or "请分析以下附件。"}]
    total = 0
    has_images = False
    for attachment in req.attachments:
        try:
            data = base64.b64decode(attachment.data, validate=True)
        except (ValueError, binascii.Error):
            raise HTTPException(400, "附件编码无效")
        total += len(data)
        if not data or len(data) > MAX_FILE_BYTES or total > MAX_TOTAL_BYTES:
            raise HTTPException(413, "每个附件须为 1 字节至 5 MB，合计不超过 15 MB")
        ext = Path(attachment.name).suffix.lower()
        try:
            if ext in IMAGE_EXTENSIONS:
                from PIL import Image
                with Image.open(io.BytesIO(data)) as img:
                    if img.width * img.height > 25_000_000:
                        raise ValueError("图片像素过多")
                    fmt = img.format
                    img.verify()
                mime = {"PNG": "image/png", "JPEG": "image/jpeg", "WEBP": "image/webp", "GIF": "image/gif"}.get(fmt)
                if not mime:
                    raise ValueError("不支持此图片格式")
                blocks.append({"type": "text", "text": "图片附件：" + attachment.name})
                blocks.append({"type": "image_url", "image_url": {"url": f"data:{mime};base64,{attachment.data}"}})
                has_images = True
                continue
            if ext in TEXT_EXTENSIONS:
                text = data.decode("utf-8-sig")
            elif ext == ".pdf":
                from pypdf import PdfReader
                reader = PdfReader(io.BytesIO(data))
                if reader.is_encrypted or len(reader.pages) > 100:
                    raise ValueError("仅支持未加密且不超过 100 页的 PDF")
                parts = []
                for page in reader.pages:
                    parts.append(page.extract_text() or "")
                    if sum(map(len, parts)) > MAX_TEXT_CHARS:
                        raise ValueError("文档文字超过 10 万字")
                text = "\n".join(parts)
            elif ext == ".docx":
                with zipfile.ZipFile(io.BytesIO(data)) as archive:
                    if sum(item.file_size for item in archive.infolist()) > 25 * 1024 * 1024:
                        raise ValueError("文档解压后过大")
                from docx import Document
                doc = Document(io.BytesIO(data))
                text = "\n".join([p.text for p in doc.paragraphs] + ["\t".join(c.text for c in row.cells) for table in doc.tables for row in table.rows])
            else:
                raise ValueError("支持 PNG/JPEG/WebP/GIF、UTF-8 文本、PDF 和 DOCX")
            if not text.strip():
                raise ValueError("未提取到文字；扫描件请转成图片后使用多模态模型")
            if len(text) > MAX_TEXT_CHARS:
                raise ValueError("文档文字超过 10 万字")
            blocks.append({"type": "text", "text": f"附件：{attachment.name}\n{text}"})
        except Exception as exc:
            raise HTTPException(400, f"无法读取附件 {attachment.name}：请检查文件格式、大小或文字编码") from exc
    return blocks, has_images


async def save_refusal(session_id, message, user_id):
    from app.agent import runtime
    from langchain_core.messages import HumanMessage, AIMessage
    agent = await runtime.get_agent(session_id, user_id=user_id, history_only=True)
    from app.agent.conversation_history import sync_checkpoints
    await sync_checkpoints(agent, session_id)
    await agent.aupdate_state({"configurable": {"thread_id": session_id}},
                             {"messages": [HumanMessage(content=message), AIMessage(content=REFUSAL)]})
    await sync_checkpoints(agent, session_id)
    runtime.invalidate_agents(session_id)
    return REFUSAL
